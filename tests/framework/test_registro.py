"""Pruebas del propio framework: el registro CSV debe ser incremental."""

import csv
import textwrap
from pathlib import Path

import pytest

from auto_testing.registro import RegistroEscenarios


RAIZ = Path(__file__).resolve().parents[2]


@pytest.fixture(autouse=True)
def _pythonpath(monkeypatch):
    monkeypatch.setenv("PYTHONPATH", str(RAIZ))


def _correr(pytester, *args):
    """Ejecuta pytest en un subproceso (aislado del plugin de Playwright)."""
    return pytester.runpytest_subprocess("-p", "no:cacheprovider", "-p", "no:playwright", *args)


def _leer(ruta):
    with ruta.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


@pytest.mark.escenario(
    modulo="Framework",
    funcionalidad="Registro CSV",
    descripcion="Asignar IDs consecutivos y no duplicar escenarios",
    resultado_esperado="ESC-0001, ESC-0002; registrar de nuevo no crea filas",
    prioridad="Alta",
)
def test_ids_incrementales(tmp_path):
    esc_csv, ej_csv = tmp_path / "e.csv", tmp_path / "x.csv"
    reg = RegistroEscenarios(esc_csv, ej_csv)
    assert reg.registrar("t::a").id == "ESC-0001"
    assert reg.registrar("t::b").id == "ESC-0002"
    reg.guardar()

    reg2 = RegistroEscenarios(esc_csv, ej_csv)
    assert reg2.registrar("t::a").id == "ESC-0001"
    assert reg2.registrar("t::c").id == "ESC-0003"
    reg2.guardar()
    assert [f["id"] for f in _leer(esc_csv)] == ["ESC-0001", "ESC-0002", "ESC-0003"]


@pytest.mark.escenario(
    modulo="Framework",
    funcionalidad="Registro CSV",
    descripcion="Corridas sucesivas agregan historial y actualizan el catálogo",
    resultado_esperado="Catálogo sin duplicados; historial crece por corrida",
    prioridad="Alta",
)
def test_corridas_sucesivas(pytester):
    pytester.makeconftest('pytest_plugins = ["auto_testing.plugin"]')
    pytester.makepyfile(
        test_demo=textwrap.dedent(
            """
            import pytest

            @pytest.mark.escenario(modulo="Demo", descripcion="pasa")
            def test_ok():
                assert True

            def test_falla():
                '''Falla a propósito'''
                assert 1 == 2
            """
        )
    )
    _correr(pytester).assert_outcomes(passed=1, failed=1)
    _correr(pytester).assert_outcomes(passed=1, failed=1)

    escenarios = _leer(pytester.path / "reportes" / "escenarios.csv")
    assert len(escenarios) == 2
    por_desc = {e["descripcion"]: e for e in escenarios}
    assert por_desc["pasa"]["ultimo_estado"] == "EXITOSO"
    assert por_desc["Falla a propósito"]["ultimo_estado"] == "FALLIDO"
    assert por_desc["Falla a propósito"]["ejecuciones"] == "2"
    assert por_desc["Falla a propósito"]["fallos"] == "2"

    ejecuciones = _leer(pytester.path / "reportes" / "ejecuciones.csv")
    assert len(ejecuciones) == 4

    # Una prueba nueva solo agrega un escenario con el siguiente ID.
    pytester.makepyfile(test_nuevo="def test_nuevo():\n    pass\n")
    _correr(pytester)
    escenarios = _leer(pytester.path / "reportes" / "escenarios.csv")
    assert [e["id"] for e in escenarios] == ["ESC-0001", "ESC-0002", "ESC-0003"]


@pytest.mark.escenario(
    modulo="Framework",
    funcionalidad="Levantamiento",
    descripcion="--collect-only levanta escenarios sin ejecutarlos",
    resultado_esperado="Catálogo con estado NO_EJECUTADO y sin historial",
)
def test_collect_only(pytester):
    pytester.makeconftest('pytest_plugins = ["auto_testing.plugin"]')
    pytester.makepyfile(test_demo="def test_a():\n    pass\n")
    _correr(pytester, "--collect-only")
    escenarios = _leer(pytester.path / "reportes" / "escenarios.csv")
    assert escenarios[0]["ultimo_estado"] == "NO_EJECUTADO"
    assert not (pytester.path / "reportes" / "ejecuciones.csv").exists()


@pytest.mark.escenario(
    modulo="Framework",
    funcionalidad="Reporte PDF",
    descripcion="Generar el reporte PDF al terminar la corrida",
    resultado_esperado="reportes/reporte.pdf existe y es un PDF válido; --sin-pdf no lo genera",
    prioridad="Alta",
)
def test_reporte_pdf(pytester):
    pytester.makeconftest('pytest_plugins = ["auto_testing.plugin"]')
    pytester.makepyfile(
        test_demo=textwrap.dedent(
            """
            import pytest

            def test_ok():
                pass

            def test_falla():
                assert False, "algo salió mal"

            @pytest.mark.xfail(reason="defecto conocido")
            def test_defecto():
                assert False
            """
        )
    )
    pdf = pytester.path / "reportes" / "reporte.pdf"
    _correr(pytester, "--sin-pdf")
    assert not pdf.exists()

    resultado = _correr(pytester)
    resultado.assert_outcomes(passed=1, failed=1, xfailed=1)
    assert pdf.read_bytes().startswith(b"%PDF")
    ejecuciones = _leer(pytester.path / "reportes" / "ejecuciones.csv")
    assert {"defecto conocido"} <= {e["mensaje"] for e in ejecuciones}


@pytest.mark.escenario(
    modulo="Framework",
    funcionalidad="Registro CSV",
    descripcion="El catálogo usa la coma solo como separador",
    resultado_esperado="Textos con comas y comillas se guardan sin comillas; cada línea tiene el mismo número de columnas",
    prioridad="Alta",
)
def test_csv_sin_comas_en_textos(tmp_path):
    esc_csv = tmp_path / "e.csv"
    reg = RegistroEscenarios(esc_csv, tmp_path / "x.csv")
    reg.registrar(
        "t::a",
        descripcion='Subtotal, impuesto y total "con IVA"',
        resultado_esperado="Subtotal 39.98, impuesto 3.20,total 43.18\nfin",
    )
    reg.guardar()
    lineas = esc_csv.read_text(encoding="utf-8").splitlines()
    assert '"' not in esc_csv.read_text(encoding="utf-8")
    columnas = {len(linea.split(",")) for linea in lineas}
    assert columnas == {len(lineas[0].split(","))}
    assert ";" not in esc_csv.read_text(encoding="utf-8")  # Excel en español separa por ';'
    fila = _leer(esc_csv)[0]
    assert fila["descripcion"] == "Subtotal - impuesto y total 'con IVA'"
    assert fila["resultado_esperado"] == "Subtotal 39.98 - impuesto 3.20 - total 43.18 fin"


@pytest.mark.escenario(
    modulo="Framework",
    funcionalidad="Registro CSV",
    descripcion="Generar el catálogo en Excel (.xlsx)",
    resultado_esperado="escenarios.xlsx tiene encabezados en español y una fila por escenario con su estado",
)
def test_catalogo_excel(tmp_path):
    from openpyxl import load_workbook

    esc_csv = tmp_path / "escenarios.csv"
    reg = RegistroEscenarios(esc_csv, tmp_path / "x.csv")
    reg.registrar("t::a", descripcion="Agregar productos, calcular total", modulo="Carrito")
    reg.registrar("t::b", descripcion="Quitar producto")
    reg.registrar_resultado("t::a", "EXITOSO", "2026-01-01T00:00:00")
    reg.guardar()
    hoja = load_workbook(tmp_path / "escenarios.xlsx").active
    filas = list(hoja.iter_rows(values_only=True))
    assert filas[0][:4] == ("ID", "Módulo", "Funcionalidad", "Descripción")
    assert [f[0] for f in filas[1:]] == ["ESC-0001", "ESC-0002"]
    assert filas[1][3] == "Agregar productos - calcular total"
    assert filas[1][7] == "EXITOSO" and filas[2][7] == "NO_EJECUTADO"
