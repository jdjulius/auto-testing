"""Exporta el catálogo de escenarios a Excel (.xlsx)."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# columna del CSV -> (encabezado, ancho)
COLUMNAS = {
    "id": ("ID", 11),
    "modulo": ("Módulo", 20),
    "funcionalidad": ("Funcionalidad", 24),
    "descripcion": ("Descripción", 55),
    "resultado_esperado": ("Resultado esperado", 60),
    "prioridad": ("Prioridad", 11),
    "navegador": ("Navegador", 12),
    "ultimo_estado": ("Último estado", 15),
    "ultima_ejecucion": ("Última ejecución", 20),
    "ejecuciones": ("Ejecuciones", 12),
    "fallos": ("Fallos", 9),
    "fecha_alta": ("Fecha de alta", 20),
    "nodeid": ("Prueba (nodeid)", 70),
}

COLORES_ESTADO = {
    "EXITOSO": "C6EFCE",
    "FALLIDO": "FFC7CE",
    "ERROR": "E4C1F9",
    "XFAIL": "FFEB9C",
    "XPASS": "BDD7EE",
    "OMITIDO": "D9D9D9",
    "NO_EJECUTADO": "F2F2F2",
}


def exportar_excel(ruta: Path, escenarios: list[dict]) -> Path:
    ruta = Path(ruta)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Escenarios"

    claves = list(COLUMNAS)
    hoja.append([COLUMNAS[c][0] for c in claves])
    for esc in escenarios:
        hoja.append([esc.get(c, "") for c in claves])

    encabezado = PatternFill("solid", start_color="132322")
    for i, clave in enumerate(claves, start=1):
        celda = hoja.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(vertical="center")
        hoja.column_dimensions[get_column_letter(i)].width = COLUMNAS[clave][1]

    col_estado = claves.index("ultimo_estado") + 1
    ajustar = {claves.index(c) + 1 for c in ("descripcion", "resultado_esperado")}
    for fila in range(2, hoja.max_row + 1):
        estado = hoja.cell(row=fila, column=col_estado)
        color = COLORES_ESTADO.get(str(estado.value))
        if color:
            estado.fill = PatternFill("solid", start_color=color)
        for col in ajustar:
            hoja.cell(row=fila, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    hoja.freeze_panes = "B2"
    hoja.auto_filter.ref = hoja.dimensions

    ruta.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=ruta.parent, prefix=".escenarios", suffix=".xlsx")
    os.close(fd)
    libro.save(tmp)
    os.replace(tmp, ruta)
    return ruta
