"""Exporta el catálogo de escenarios a Excel (.xlsx)."""

from __future__ import annotations

import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# columna del CSV -> (encabezado, ancho)
COLUMNAS = {
    "id": ("ID", 11),
    "capa": ("Capa", 11),
    "modulo": ("Módulo", 20),
    "funcionalidad": ("Funcionalidad", 24),
    "descripcion": ("Descripción", 55),
    "resultado_esperado": ("Resultado esperado", 60),
    "prioridad": ("Prioridad", 11),
    "navegador": ("Navegador", 12),
    "ultimo_estado": ("Último estado", 15),
    "ultima_ejecucion": ("Última ejecución", 20),
    "ejecuciones": ("Ejecuciones", 12),
    "fallos": ("Fallos", 9),
    "fecha_alta": ("Fecha de alta", 20),
    "nodeid": ("Prueba (nodeid)", 70),
}

COLORES_ESTADO = {
    "EXITOSO": "C6EFCE",
    "FALLIDO": "FFC7CE",
    "ERROR": "E4C1F9",
    "XFAIL": "FFEB9C",
    "XPASS": "BDD7EE",
    "OMITIDO": "D9D9D9",
    "NO_EJECUTADO": "F2F2F2",
}


def exportar_excel(ruta: Path, escenarios: list[dict]) -> Path:
    """Catálogo de escenarios."""
    return exportar_tabla(
        ruta, escenarios, COLUMNAS, hoja="Escenarios",
        colores={"ultimo_estado": COLORES_ESTADO},
        ajustar=("descripcion", "resultado_esperado"),
    )


def exportar_tabla(
    ruta: Path,
    filas: list[dict],
    columnas: dict[str, tuple[str, int]],
    *,
    hoja: str = "Hoja1",
    colores: dict[str, dict[str, str]] | None = None,
    ajustar: tuple[str, ...] = (),
) -> Path:
    """Escribe `filas` en un .xlsx con encabezado, anchos, filtros y fila fija.

    `columnas` es {clave: (encabezado, ancho)}; `colores` pinta una columna según su
    valor ({clave: {valor: color}}); `ajustar` son las columnas con texto en varias líneas.
    """
    ruta = Path(ruta)
    libro = Workbook()
    hoja_xl = libro.active
    hoja_xl.title = hoja

    claves = list(columnas)
    hoja_xl.append([columnas[c][0] for c in claves])
    for fila in filas:
        hoja_xl.append([fila.get(c, "") for c in claves])

    encabezado = PatternFill("solid", start_color="132322")
    for i, clave in enumerate(claves, start=1):
        celda = hoja_xl.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = encabezado
        celda.alignment = Alignment(vertical="center")
        hoja_xl.column_dimensions[get_column_letter(i)].width = columnas[clave][1]

    pintar = {claves.index(c) + 1: paleta for c, paleta in (colores or {}).items() if c in claves}
    envolver = {claves.index(c) + 1 for c in ajustar if c in claves}
    for fila in range(2, hoja_xl.max_row + 1):
        for col, paleta in pintar.items():
            celda = hoja_xl.cell(row=fila, column=col)
            color = paleta.get(str(celda.value))
            if color:
                celda.fill = PatternFill("solid", start_color=color)
        for col in envolver:
            hoja_xl.cell(row=fila, column=col).alignment = Alignment(wrap_text=True, vertical="top")

    hoja_xl.freeze_panes = "B2"
    hoja_xl.auto_filter.ref = hoja_xl.dimensions

    ruta.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=ruta.parent, prefix=f".{ruta.stem}", suffix=".xlsx")
    os.close(fd)
    libro.save(tmp)
    os.replace(tmp, ruta)
    return ruta
